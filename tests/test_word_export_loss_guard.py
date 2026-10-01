"""Regression for two observed converter losses, not OCR accuracy."""
import json
import pytest
from test_saved_word_export import native, converter_available
from mp4_analysis.thin.output import export_saved_word
from mp4_analysis.thin.utils import file_hash


def test_html_cell_is_not_markdown_list(native, tmp_path):
    converter_available()
    from docx import Document
    (native/'page.md').write_text('<div><table><tr><td>+ I</td><td>4</td></tr></table></div>')
    cache = json.loads((native/'adapter.json').read_text())
    cache['files']['page.md'] = file_hash(native/'page.md')
    (native/'adapter.json').write_text(json.dumps(cache))
    export_saved_word(native, tmp_path/'derived')
    assert Document(tmp_path/'derived/document.docx').tables[0].cell(0, 0).text == '+ I'


def test_overflow_cell_loss_is_not_published(native, tmp_path):
    converter_available()
    markup = '<table><tr><td rowspan="2">A</td><td>B</td></tr><tr><td>C</td><td>DO NOT DROP</td></tr></table>'
    (native/'page.md').write_text(markup)
    cache = json.loads((native/'adapter.json').read_text())
    cache['files']['page.md'] = file_hash(native/'page.md')
    (native/'adapter.json').write_text(json.dumps(cache))
    with pytest.raises(ValueError, match='table count/cell text changed'):
        export_saved_word(native, tmp_path/'derived')
    assert not (tmp_path/'derived').exists()
