"""
exporters/xlsx.py - 工业级 Excel 导出器
将 Document IR 中的表格数据结构化输出为标准 .xlsx，保留十六进制纯文本格式与单元格样式
"""

import os
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ExcelExporter:
    def __init__(self):
        self.font_header = Font(name="Microsoft YaHei", size=10, bold=True)
        self.font_data = Font(name="Microsoft YaHei", size=9)
        self.fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        self.thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )
        self.align_center = Alignment(horizontal="center", vertical="center")
        self.align_left = Alignment(horizontal="left", vertical="center")

    def export_table(self, title: str, headers: List[str], rows: List[List[Any]], output_path: str):
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:30] if title else "Sheet1"
        ws.views.sheetView[0].showGridLines = True

        # 写入表头
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=str(h))
            cell.font = self.font_header
            cell.fill = self.fill_header
            cell.alignment = self.align_center
            cell.border = self.thin_border

        # 写入数据行
        for r_idx, row in enumerate(rows, 2):
            for c_idx, val in enumerate(row, 1):
                # 强制以纯文本存储，防止 16 进制或位宽格式被 Excel 破坏
                cell = ws.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else "")
                cell.font = self.font_data
                cell.border = self.thin_border
                cell.alignment = self.align_center if c_idx == 1 or len(str(val)) < 10 else self.align_left

        # 自动调整列宽
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 4, 50))

        wb.save(output_path)
        return output_path
