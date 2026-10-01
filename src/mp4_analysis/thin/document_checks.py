"""Read-only export contracts, not OCR or an inferred table-grid solver.

Only CR/LF and HTML nonbreaking-space representation are normalized. Word
paragraph boundaries, word separators, image order and declared cell spans
remain significant. These checks do not certify source accuracy or rendering.
"""
from __future__ import annotations

import hashlib
import io
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from .utils import file_hash

_NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
       'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
       'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def literal_text(value):
    """Preserve spaces/tabs and their multiplicity; never compact identifiers."""
    return value.replace('\r\n', '\n').replace('\r', '\n').replace('\u00a0', ' ')


def declared_cells(table):
    """Evaluate explicit HTML spans for verification only; never infer/fix them.

    Coordinates are 1-based. Ambiguous overlaps or unreasonable expanded area
    are rejected rather than repaired. This function never writes a worksheet.
    """
    cells = {c['id']: c for c in table['cells']}
    if len(cells) != len(table['cells']):
        raise ValueError('Duplicate declared cell ID')
    occupied, anchors, merges, seen = set(), {}, [], set()
    height, width = len(table['rows']), 0
    for row_index, row in enumerate(table['rows'], 1):
        column = 1
        for cid in row:
            if cid not in cells or cid in seen:
                raise ValueError('Missing or repeated declared cell')
            seen.add(cid)
            cell = cells[cid]
            rs, cs = cell['rowspan'], cell['colspan']
            if type(rs) is not int or type(cs) is not int or min(rs, cs) < 1:
                raise ValueError('Invalid explicit table span')
            while (row_index, column) in occupied:
                column += 1
            if row_index + rs - 1 > height or column + cs - 1 > 16384:
                raise ValueError('Explicit span exceeds declared rows or worksheet limit')
            if len(occupied) + rs * cs > 200000:
                raise ValueError('Declared table expansion exceeds verification limit')
            region = {(r, c) for r in range(row_index, row_index + rs)
                      for c in range(column, column + cs)}
            if occupied.intersection(region):
                raise ValueError('Overlapping explicit spans; no automatic repair')
            occupied.update(region)
            anchors[row_index, column] = cell['text']
            if rs > 1 or cs > 1:
                merges.append((row_index, column, row_index + rs - 1, column + cs - 1))
            width = max(width, column + cs - 1)
            column += cs
    if seen != set(cells):
        raise ValueError('Unreferenced declared cell')
    return anchors, sorted(merges), (height, width)


def check_worksheet(worksheet, table):
    """Compare candidate -> converted/saved sheet, including empty-cell spans."""
    from openpyxl.cell.cell import MergedCell
    anchors, merges, shape = declared_cells(table)
    actual_shape = (worksheet.max_row, worksheet.max_column)
    actual_merges = sorted((m.min_row, m.min_col, m.max_row, m.max_col)
                           for m in worksheet.merged_cells.ranges)
    values = {(c.row, c.column): (c.value, c.data_type)
              for row in worksheet for c in row if not isinstance(c, MergedCell)}
    wrong = []
    for coordinate, text in anchors.items():
        value, kind = values.get(coordinate, (None, None))
        same = (text == '' and value in (None, '')) or (
            isinstance(value, str) and literal_text(value) == literal_text(text)
            and kind in {'s', 'inlineStr'})
        if not same:
            wrong.append(list(coordinate))
    extra = [list(rc) for rc, (value, _) in values.items()
             if rc not in anchors and value not in (None, '')]
    same = not wrong and not extra and merges == actual_merges and shape == actual_shape
    return {'status': 'CONSISTENT' if same else 'EXPORT_MISMATCH',
            'literal_cells_equal': not wrong and not extra,
            'declared_spans_equal': merges == actual_merges,
            'declared_extent_equal': shape == actual_shape,
            'mismatched_coordinates': wrong, 'unexpected_coordinates': extra,
            'expected_merges': merges, 'actual_merges': actual_merges,
            'expected_shape': shape, 'actual_shape': actual_shape,
            'source_geometry_verified': False}


def _paragraph_text(node):
    parts = []
    for child in node.iter():
        if child.tag == '{'+_NS['w']+'}t':
            parts.append(child.text or '')
        elif child.tag == '{'+_NS['w']+'}tab':
            parts.append('\t')
        elif child.tag in {'{'+_NS['w']+'}br', '{'+_NS['w']+'}cr'}:
            if child.get('{'+_NS['w']+'}type', 'textWrapping') == 'textWrapping':
                parts.append('\n')
    return literal_text(''.join(parts))


def _pixels(data):
    from PIL import Image
    with Image.open(io.BytesIO(data)) as image:
        image.load()
        rgb = image.convert('RGBA')
        return hashlib.sha256(str(rgb.size).encode() + rgb.tobytes()).hexdigest()


def check_docx(path, document):
    """Compare ordered text/table/image events with the frozen candidate.

    Pictures are checked by embedded resource pixels (metadata-only re-encoding
    is allowed). Missing/external/broken image relationships fail closed.
    """
    expected, observed, errors = [], [], []
    root_directory = Path(path).parent.resolve()
    for unit in document['units']:
        for block in unit['blocks']:
            if block.get('excluded'):
                continue
            if block['kind'] == 'image':
                resource = (root_directory / block['image']).resolve()
                if not resource.is_relative_to(root_directory):
                    raise ValueError('Candidate image outside bundle')
                expected.append(('image', _pixels(resource.read_bytes())))
            elif block['kind'] == 'table':
                table = block['table']
                if table.get('caption'):
                    expected.append(('text', literal_text(table['caption'])))
                cells = {c['id']: c['text'] for c in table['cells']}
                expected.append(('table', [literal_text(cells[cid])
                    for row in table['rows'] for cid in row]))
            else:
                expected.append(('text', literal_text(block.get('text', ''))))
    with zipfile.ZipFile(path) as archive:
        body = ET.fromstring(archive.read('word/document.xml')).find('w:body', _NS)
        if body is None:
            raise ValueError('DOCX has no document body')
        relationships = {}
        if 'word/_rels/document.xml.rels' in archive.namelist():
            relationships = {r.get('Id'): r for r in ET.fromstring(
                archive.read('word/_rels/document.xml.rels'))}
        for element in body:
            if element.tag == '{'+_NS['w']+'}p':
                pictures = element.findall('.//a:blip', _NS)
                text = _paragraph_text(element)
                if not pictures or text:
                    observed.append(('text', text))
                for picture in pictures:
                    rel = relationships.get(picture.get('{'+_NS['r']+'}embed'))
                    try:
                        if (rel is None or rel.get('TargetMode') == 'External'
                                or not rel.get('Type', '').endswith('/image')
                                or picture.get('{'+_NS['r']+'}link')):
                            raise ValueError('Missing or external image relationship')
                        target = rel.get('Target', '')
                        member = posixpath.normpath(posixpath.join('word', target))
                        if ':' in target or '\\' in target or not member.startswith('word/media/'):
                            raise ValueError('Unexpected embedded image path')
                        observed.append(('image', _pixels(archive.read(member))))
                    except (ValueError, OSError, KeyError) as exc:
                        errors.append(str(exc))
                        observed.append(('image', 'INVALID_RESOURCE'))
            elif element.tag == '{'+_NS['w']+'}tbl':
                cells = []
                for row in element.findall('./w:tr', _NS):
                    for cell in row.findall('./w:tc', _NS):
                        merge = cell.find('./w:tcPr/w:vMerge', _NS)
                        if merge is not None and merge.get('{'+_NS['w']+'}val') != 'restart':
                            continue
                        cells.append('\n'.join(_paragraph_text(p) for p in cell.findall('./w:p', _NS)))
                observed.append(('table', cells))
                if element.findall('.//a:blip', _NS):
                    errors.append('Unexpected picture inside a text-only candidate table')
            elif element.tag != '{'+_NS['w']+'}sectPr':
                errors.append('Unsupported Word body element: '+element.tag)
    same = observed == expected and not errors
    return {'status': 'CONSISTENT' if same else 'EXPORT_MISMATCH',
            'ordered_literal_blocks_equal': observed == expected,
            'images_equal': [v for k, v in expected if k == 'image'] ==
                            [v for k, v in observed if k == 'image'] and not errors,
            'expected_image_count': sum(k == 'image' for k, _ in expected),
            'actual_image_count': sum(k == 'image' for k, _ in observed),
            'expected_block_count': len(expected), 'actual_block_count': len(observed),
            'resource_errors': errors, 'file_sha256': file_hash(path),
            'whitespace_policy': 'CRLF_to_LF_NBSP_to_space; preserve separators and block boundaries',
            'source_accuracy_verified': False, 'layout_rendered': False,
            'merge_geometry_verified': False}
