"""Output-only regressions: real external converter; no OCR or cell repairs."""
import json
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from mp4_analysis.thin.output import export_saved_word
from mp4_analysis.thin.utils import file_hash

TABLE = ('<table><tr><td rowspan="2">logo</td><td>ET6601</td>'
         '<td rowspan="2">5/16</td><td rowspan="2">V1.0</td></tr>'
         '<tr><td>SARC</td></tr><tr><td colspan="4">0000 0x0000_0000 0x0000_7FFF</td></tr></table>')


@pytest.fixture
def native(tmp_path):
    root = tmp_path/'native'
    root.mkdir()
    (root/'page.md').write_text('# Heading\n\n'+TABLE+'\n\nADC数字控制器不能早于模拟ADC解复位\n', encoding='utf-8')
    (root/'adapter.json').write_text(json.dumps({'errors': [], 'signature': {'parser': 'fixture'},
        'files': {'page.md': file_hash(root/'page.md')}}), encoding='utf-8')
    return root


def converter_available():
    pytest.importorskip('docx')
    executable = shutil.which('pandoc')
    if not executable or subprocess.check_output([executable, '--version']).decode().splitlines()[0] != 'pandoc 3.1.11.1':
        pytest.skip('real export test requires external Pandoc 3.1.11.1')


def test_real_spans_values_and_native_bytes(native, tmp_path):
    converter_available()
    from docx import Document
    before = {p.name: p.read_bytes() for p in native.iterdir()}
    out = tmp_path/'derived'
    info = export_saved_word(native, out)
    document = Document(out/'document.docx')
    table = document.tables[0]
    assert len(table.rows) == 3 and len(table.columns) == 4
    assert table.cell(1, 1).text == 'SARC'
    assert table.cell(0, 0)._tc is table.cell(1, 0)._tc
    assert table.cell(0, 2)._tc is table.cell(1, 2)._tc
    assert table.cell(0, 3)._tc is table.cell(1, 3)._tc
    assert table.cell(2, 0)._tc is table.cell(2, 3)._tc
    assert table.cell(2, 0).text == '0000 0x0000_0000 0x0000_7FFF'
    assert 'ADC数字控制器不能早于模拟ADC解复位' in '\n'.join(p.text for p in document.paragraphs)
    assert before == {p.name: p.read_bytes() for p in native.iterdir()}
    assert info['inference_performed'] is False and info['accuracy_verified'] is False


def test_real_image_bytes_retained(native, tmp_path):
    converter_available()
    from PIL import Image
    Image.new('RGB', (20, 20), 'white').save(native/'image.png')
    (native/'page.md').write_text('<img src="image.png" />', encoding='utf-8')
    cache = json.loads((native/'adapter.json').read_text())
    cache['files'] = {n: file_hash(native/n) for n in ('image.png', 'page.md')}
    (native/'adapter.json').write_text(json.dumps(cache))
    out = tmp_path/'derived'
    export_saved_word(native, out)
    with zipfile.ZipFile(out/'document.docx') as z:
        assert (native/'image.png').read_bytes() in [z.read(n) for n in z.namelist() if n.startswith('word/media/')]


@pytest.mark.parametrize('change', ['corrupt', 'missing', 'errors', 'traversal', 'empty', 'two_markdown'])
def test_invalid_native_cache_rejected(native, tmp_path, change):
    cache = json.loads((native/'adapter.json').read_text())
    if change == 'corrupt': (native/'page.md').write_text('changed')
    elif change == 'missing': (native/'page.md').unlink()
    elif change == 'errors': cache['errors'] = ['failed export']
    elif change == 'traversal': cache['files'] = {'../elsewhere.md': 'x'}
    elif change == 'empty': cache['files'] = {}
    else:
        (native/'other.md').write_text('other')
        cache['files']['other.md'] = file_hash(native/'other.md')
    (native/'adapter.json').write_text(json.dumps(cache))
    with pytest.raises(ValueError): export_saved_word(native, tmp_path/'derived')
    assert not (tmp_path/'derived').exists()


@pytest.mark.parametrize('target', ['same', 'child', 'parent', 'existing'])
def test_no_overwrite(native, tmp_path, target):
    out = {'same':native, 'child':native/'out', 'parent':tmp_path, 'existing':tmp_path/'existing'}[target]
    if target == 'existing': out.mkdir()
    with pytest.raises((ValueError, FileExistsError)): export_saved_word(native, out)


@pytest.mark.parametrize('markup', ['<img src="https://example.invalid/x.png">', '<img src="../secret.png">', '<script>alert(1)</script>'])
def test_external_or_active_resources_fail_closed(native, tmp_path, markup):
    converter_available()
    (native/'page.md').write_text(markup)
    cache = json.loads((native/'adapter.json').read_text())
    cache['files']['page.md'] = file_hash(native/'page.md')
    (native/'adapter.json').write_text(json.dumps(cache))
    with pytest.raises(ValueError): export_saved_word(native, tmp_path/'derived')
    assert not (tmp_path/'derived').exists()


def test_missing_converter_is_explicit(native, tmp_path):
    with pytest.raises(RuntimeError, match='install external Pandoc'):
        export_saved_word(native, tmp_path/'derived', pandoc='missing-pandoc-for-test')


def test_timeout_does_not_publish(native, tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, 'which', lambda _: '/fake/pandoc')
    def fail(*args, **kwargs):
        raise subprocess.TimeoutExpired('pandoc', 0.01)
    monkeypatch.setattr(subprocess, 'run', fail)
    with pytest.raises(subprocess.TimeoutExpired): export_saved_word(native, tmp_path/'derived')
    assert not (tmp_path/'derived').exists()
