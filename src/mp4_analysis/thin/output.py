"""Small source index, integrity checks and links to UNMODIFIED native outputs."""
from __future__ import annotations

import html
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET


def inspect_workbook(path):
    """Read OOXML only; do not evaluate formulas or convert model confidence to accuracy."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        sheets = []
        for name in archive.namelist():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                root = ET.fromstring(archive.read(name))
                sheets.append({'part': name, 'rows':len(root.findall('.//s:row', ns)),
                               'cells':len(root.findall('.//s:c', ns)),
                               'formulas':len(root.findall('.//s:f', ns)),
                               'merges':len(root.findall('.//s:mergeCell', ns))})
    return {'sheets':sheets, 'formula_review_required':any(s['formulas'] for s in sheets)}


def write_index(output, manifest, items, *, stitch=None, reconstruction=None):
    output = Path(output)
    errors = [item for item in items if item.get('error') or item.get('native', {}).get('errors')]
    expected = reconstruction['parser_inputs'] if reconstruction else len(manifest.get('frames', items))
    pending = max(0, expected-len(items))
    report = {'status':'NO_CONTENT' if not items else 'PARTIAL_FAILURE' if errors or pending else 'REVIEW_REQUIRED',
              'source':manifest, 'parser_attempts':len(items),
              'expected_parser_inputs':expected, 'pending_parser_inputs':pending,
              'cache_hits':sum(item.get('native', {}).get('cache_hit', False) for item in items),
              'items':items, 'stitch_preview':stitch, 'reconstruction':reconstruction,
              'accuracy_verified':False, 'content_completeness_verified':False,
              'limitations':['Frames are observations, not original document pages.',
                 'No semantic rewriting or text-similarity deletion.',
                 'SCANS composites, when requested, remain unverified; no cross-batch content deduplication.',
                 'Sampled or capped runs cannot establish coverage.',
                 'Native model outputs, including table row/column assignments, require review.']}
    (output/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    lines = ['# 录屏资料提取索引（待核验）', '',
             '原生结果按画面保存；帧数不是原文页数。实验拼接只生成待核验候选，不代表图表完整恢复。', '']
    cards = []
    for item in items:
        frame = item['frame']
        title = f"画面 {frame['frame_index']} · {frame['start_time']:.3f} 秒"
        lines += [f'## {title}', '']
        links=[]
        for name in item.get('native', {}).get('files', {}):
            if Path(name).suffix.lower() in {'.md','.html','.xlsx','.docx','.json'}:
                relative = f"{item['directory']}/{name}"
                lines.append(f'- [{name}](<{relative}>)')
                links.append(f'<a href="{html.escape(relative,quote=True)}">{html.escape(name)}</a>')
        source_links = []
        for source in item.get('source_frames', [frame]):
            source_links.append(f'<a href="frames/{html.escape(source["image"],quote=True)}">{source["frame_index"]}</a>')
        preview = item.get('input_image', f'frames/{frame["image"]}')
        if item.get('kind') == 'unverified_composite':
            links.append('实验拼接：几何及文字完整性未核验')
        issue = item.get('error') or item.get('native', {}).get('errors')
        if issue:
            lines.append(f'错误：{issue}')
        cards.append(f'<section><h2>{html.escape(title)}</h2><p>{" · ".join(links)}</p>'
                     f'<p>{html.escape(str(issue or "待对照原画面核验"))}</p>'
                     f'<p>源帧：{" · ".join(source_links)}</p>'
                     f'<img loading="lazy" src="{html.escape(preview,quote=True)}" alt="解析输入（源帧或未核验拼接）"></section>')
    (output/'index.md').write_text('\n'.join(lines),encoding='utf-8')
    (output/'index.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
        '<title>录屏资料提取 · 待核验</title><style>body{max-width:1100px;margin:24px auto;font:16px sans-serif;line-height:1.6}img{max-width:100%}section{border-top:1px solid #ccc;padding:12px}</style>'
        '<h1>录屏资料提取 · 待核验</h1><p>保留原画面及引擎原生结果；本索引不代表完整性或准确率验收。</p>'
        +''.join(cards)+'</html>',encoding='utf-8')
    return report


def export_saved_word(native_directory, target_directory, *, pandoc='pandoc', timeout=60):
    """Re-export ONE verified native Markdown cache without inference or mutation.

    Pandoc preserves HTML table spans through Markdown -> HTML -> DOCX. This
    opt-in output adapter never replaces the engine's DOCX or marks content as
    accepted. Pandoc 3.1.11.1 is an external executable, not a parser backend.
    """
    import math
    import shutil
    import subprocess
    import tempfile
    import time
    from copy import deepcopy
    from html.parser import HTMLParser
    from urllib.parse import unquote, urlsplit
    from .video import file_hash

    source, target = Path(native_directory).resolve(), Path(target_directory).resolve()
    if source == target or source in target.parents or target in source.parents:
        raise ValueError('native and derived output directories must be disjoint')
    if target.exists():
        raise FileExistsError('derived output exists; use a new directory')
    if isinstance(timeout, bool) or not math.isfinite(timeout) or timeout <= 0:
        raise ValueError('timeout must be finite and positive')
    cache_path = source/'adapter.json'
    cache = json.loads(cache_path.read_text(encoding='utf-8'))
    files = cache.get('files', {})
    if cache.get('errors') or not files:
        raise ValueError('native cache is not a successful export')

    def checked_path(name):
        path = (source/name).resolve()
        if not path.is_relative_to(source) or not path.is_file():
            raise ValueError(f'missing or nonlocal native resource: {name}')
        return path

    for name, digest in files.items():
        if file_hash(checked_path(name)) != digest:
            raise ValueError(f'native file hash mismatch: {name}')
    markdown = [name for name in files if Path(name).suffix.lower() == '.md']
    if len(markdown) != 1:
        raise ValueError('expected exactly one native Markdown file')
    executable = shutil.which(pandoc)
    if not executable:
        raise RuntimeError('install external Pandoc 3.1.11.1 for this opt-in export')
    began = time.monotonic()

    def invoke(arguments):
        result = subprocess.run([executable, *arguments], cwd=source, capture_output=True,
                                check=True, timeout=timeout)
        # Missing images and parse warnings must not become a silent success.
        if result.stderr.strip():
            raise RuntimeError(result.stderr.decode('utf-8', errors='replace'))
        return result.stdout

    version = invoke(['--version']).decode().splitlines()[0]
    if version != 'pandoc 3.1.11.1':
        raise RuntimeError(f'unvalidated converter version: {version}')
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='word-export-', dir=target.parent) as temporary:
        work = Path(temporary)
        intermediate = work/'intermediate.html'
        # A fragment avoids adding an artificial title. Disable smart punctuation
        # and YAML metadata so addresses/quotes and metadata are not reinterpreted.
        invoke([str(checked_path(markdown[0])), '--from=markdown+raw_html-smart-yaml_metadata_block-markdown_in_html_blocks',
                '--to=html', '-o', str(intermediate)])

        class ResourceCheck(HTMLParser):
            def __init__(self):
                super().__init__()
                self.tables, self.table, self.cell = [], None, None

            def handle_starttag(self, tag, attrs):
                if tag == 'table':
                    if self.table is not None:
                        raise ValueError('nested HTML tables require separate review')
                    self.table = []
                elif tag in {'td', 'th'}:
                    if self.cell is not None:
                        raise ValueError('malformed HTML cell nesting')
                    self.cell = []
                if tag in {'script', 'iframe', 'object', 'embed', 'link', 'base'}:
                    raise ValueError(f'unsupported active HTML element: {tag}')
                for key, value in attrs:
                    if key in {'src', 'srcset', 'data', 'poster'} and value:
                        parts = urlsplit(value)
                        if key != 'src' or parts.scheme or parts.netloc or parts.query:
                            raise ValueError('only local, hashed image resources are supported')
                        name = unquote(parts.path)
                        checked_path(name)
                        if name not in files:
                            raise ValueError(f'image is outside native cache manifest: {name}')

            def handle_data(self, data):
                if self.cell is not None:
                    self.cell.append(data)

            def handle_endtag(self, tag):
                if tag in {'td', 'th'} and self.cell is not None:
                    text = ''.join(''.join(self.cell).split())
                    if self.table is not None and text:
                        self.table.append(text)
                    self.cell = None
                elif tag == 'table' and self.table is not None:
                    self.tables.append(self.table)
                    self.table = None

        markup = ResourceCheck()
        markup.feed(intermediate.read_text(encoding='utf-8'))
        markup.close()
        # Reference styles only: no custom table layout/merge solver or cell edits.
        from docx import Document
        reference = work/'reference.docx'
        reference.write_bytes(invoke(['--print-default-data-file=reference.docx']))
        document = Document(reference)
        document.styles.element.append(deepcopy(Document().styles['Table Grid'].element))
        document.styles['Table'].base_style = document.styles['Table Grid']
        document.save(reference)
        docx = work/'document.docx'
        invoke([str(intermediate), '--from=html', '--to=docx', '--resource-path', str(source),
                '--reference-doc', str(reference), '-o', str(docx)])
        if not docx.is_file() or not zipfile.is_zipfile(docx):
            raise RuntimeError('converter did not create a valid DOCX archive')
        # Compare physical cell text, retaining order and duplicate values. This
        # is a loss alarm, NOT a row/column solver or source correctness verdict.
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        with zipfile.ZipFile(docx) as archive:
            root = ET.fromstring(archive.read('word/document.xml'))
        tables = []
        for table in root.findall('.//w:tbl', ns):
            cells = []
            for cell in table.findall('./w:tr/w:tc', ns):
                text = ''.join(''.join(t.text or '' for t in cell.findall('.//w:t', ns)).split())
                if text:
                    cells.append(text)
            tables.append(cells)
        if tables != markup.tables:
            raise ValueError('table count/cell text changed during export; native evidence retained')
        # Check again before publication; do not publish results of a changing cache.
        if any(file_hash(checked_path(n)) != h for n, h in files.items()):
            raise ValueError('native cache changed during export')
        info = {'status': 'REVIEW_REQUIRED', 'converter': version,
                'native_directory': str(source), 'native_adapter_sha256': file_hash(cache_path),
                'native_signature': cache.get('signature'), 'native_files': files,
                'source_markdown': markdown[0], 'inference_performed': False,
                'native_outputs_modified': False, 'accuracy_verified': False,
                'content_completeness_verified': False,
                'elapsed_seconds': round(time.monotonic()-began, 4),
                'files': {p.name: file_hash(p) for p in (intermediate, reference, docx)},
                'limitations': ['Reflowed copy, not original page geometry.',
                    'Existing OCR errors, watermark text and incomplete images remain.',
                    'Table spans preserved by converter do not prove source-table correctness.']}
        (work/'export.json').write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding='utf-8')
        shutil.copytree(work, target)  # refuses existing paths; leaves native evidence untouched
    return info
