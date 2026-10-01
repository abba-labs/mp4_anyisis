"""Regression definitions for the local test machine. No OCR or model calls.

Run locally with pytest. These tests were authored, not executed, by ChatGPT.
Synthetic fixtures here are not source-content accuracy evidence.
"""
import copy
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image

from mp4_analysis.thin.document_format import table_model
from mp4_analysis.thin.document_checks import check_docx, check_worksheet, literal_text
from mp4_analysis.thin.document_review import _merge_coverage
from mp4_analysis.thin.document_bundle import build_document, require_delivery_target
from mp4_analysis.thin.utils import write_json


@pytest.mark.parametrize('tag', ['sup', 'sub', 's', 'strike', 'del'])
def test_R1_semantic_table_format_is_never_flattened(tag):
    with pytest.raises(ValueError, match='cannot be flattened'):
        table_model(f'<table><tr><td>10<{tag}>3</{tag}></td></tr></table>', 'b')


def test_R1_inline_semantic_style_is_rejected():
    with pytest.raises(ValueError, match='Styled'):
        table_model('<table><tr><td><span style="vertical-align:super">3</span></td></tr></table>', 'b')


def candidate(blocks):
    return {'units': [{'id': 'u1', 'blocks': blocks}]}


def test_R2_literal_separator_policy():
    assert literal_text('A B') != literal_text('AB')
    assert literal_text('A  B') != literal_text('A B')
    assert literal_text('A\tB') != literal_text('A B')
    assert literal_text('A\u00a0B\r\nC') == 'A B\nC'


def test_R2_docx_detects_missing_word_separator(tmp_path):
    from docx import Document
    frozen = candidate([{'id':'b1', 'kind':'text', 'text':'Flash Power Switch'}])
    path = tmp_path/'document.docx'
    doc = Document()
    doc.add_paragraph('Flash Power Switch')
    doc.save(path)
    assert check_docx(path, frozen)['status'] == 'CONSISTENT'
    doc.paragraphs[0].text = 'FlashPowerSwitch'
    doc.save(path)
    assert check_docx(path, frozen)['status'] == 'EXPORT_MISMATCH'


def test_R2_docx_checks_block_boundaries(tmp_path):
    from docx import Document
    frozen = candidate([{'id':'a', 'kind':'text', 'text':'one'},
                        {'id':'b', 'kind':'text', 'text':'two'}])
    path = tmp_path/'document.docx'
    doc = Document()
    doc.add_paragraph('onetwo')
    doc.save(path)
    assert check_docx(path, frozen)['status'] == 'EXPORT_MISMATCH'


def test_R2_docx_missing_or_wrong_picture_fails(tmp_path):
    from docx import Document
    image = tmp_path/'image.png'
    other = tmp_path/'other.png'
    Image.new('RGB', (24, 24), 'red').save(image)
    Image.new('RGB', (24, 24), 'blue').save(other)
    frozen = candidate([{'id':'b1', 'kind':'image', 'image':'image.png'}])
    path = tmp_path/'document.docx'
    doc = Document()
    doc.add_picture(str(image))
    doc.save(path)
    assert check_docx(path, frozen)['status'] == 'CONSISTENT'
    missing = Document()
    missing.save(path)
    assert not check_docx(path, frozen)['images_equal']
    wrong = Document()
    wrong.add_picture(str(other))
    wrong.save(path)
    assert check_docx(path, frozen)['status'] == 'EXPORT_MISMATCH'


def test_R2_candidate_to_excel_checks_empty_merge_and_positions():
    from openpyxl import Workbook
    table = table_model('<table><tr><td colspan="2"></td><td>A B</td></tr>'
                        '<tr><td>1</td><td>2</td><td>3</td></tr></table>', 'b')
    wb = Workbook()
    ws = wb.active
    ws.merge_cells('A1:B1')
    ws['C1'] = 'A B'
    for col, value in enumerate(['1', '2', '3'], 1):
        ws.cell(2, col, value)
    assert check_worksheet(ws, table)['status'] == 'CONSISTENT'
    ws.unmerge_cells('A1:B1')
    result = check_worksheet(ws, table)
    assert not result['declared_spans_equal']
    assert result['status'] == 'EXPORT_MISMATCH'
    ws.merge_cells('A1:B1')
    ws['C1'] = 'AB'
    assert not check_worksheet(ws, table)['literal_cells_equal']
    wb.close()


def review_fixture():
    return {'units':[{'id':'u1', 'blocks':[{'id':'b1','kind':'text','text':'a'}]},
                     {'id':'u2', 'blocks':[{'id':'b2','kind':'text','text':'b'}]}],
            'evidence':{'e1':{'unit_id':'u1','sha256':'a'*64},
                        'e2':{'unit_id':'u2','sha256':'b'*64}},
            'unresolved':[{'type':'critical_source_acceptance_required','issue_id':'c1'}]}


def test_R4_incremental_coverage_accumulates_without_stale_pending():
    base = review_fixture()
    first = copy.deepcopy(base)
    _merge_coverage(base, first, ['e1'], '1'*64)
    assert first['reviewed_evidence_ids_reported'] == ['e1']
    assert [x['evidence_id'] for x in first['unresolved'] if x['type']=='source_review_pending'] == ['e2']
    second = copy.deepcopy(first)
    _merge_coverage(first, second, ['e2'], '2'*64)
    assert second['reviewed_evidence_ids_reported'] == ['e1','e2']
    assert not [x for x in second['unresolved'] if x['type']=='source_review_pending']
    assert any(x['type']=='critical_source_acceptance_required' for x in second['unresolved'])
    assert second['accuracy_verified'] is False
    third = copy.deepcopy(second)
    _merge_coverage(second, third, [], '3'*64)
    assert third['reviewed_evidence_ids_reported'] == ['e1','e2']


def test_R4_changed_unit_invalidates_only_its_old_coverage():
    base = review_fixture()
    first = copy.deepcopy(base)
    _merge_coverage(base, first, ['e1','e2'], '1'*64)
    revised = copy.deepcopy(first)
    revised['units'][0]['blocks'][0]['text'] = 'changed'
    _merge_coverage(first, revised, [], '2'*64)
    assert revised['reviewed_evidence_ids_reported'] == ['e2']
    assert [x['evidence_id'] for x in revised['unresolved'] if x['type']=='source_review_pending'] == ['e1']


def test_R3_missing_group_does_not_block_available_group(tmp_path):
    from mp4_analysis.thin.batch_cli import load_plan, _prepare
    available = tmp_path/'available'
    available.mkdir()
    Image.new('RGB', (24, 24), 'white').save(available/'001.png')
    plan_path = tmp_path/'plan.json'
    write_json(plan_path, {'schema':1, 'output_root':'out', 'documents':[
        {'id':'missing','source':'missing','full_image':True},
        {'id':'available','source':'available','full_image':True}]})
    plan = load_plan(plan_path)  # Must not reject ordinary input unavailability.
    root = Path(plan['output_root'])
    root.mkdir()
    prepared = _prepare(plan, root)
    assert [d['status'] for d in prepared['documents']] == ['PREPARE_FAILED', 'PREPARED']
    assert prepared['documents'][1]['ready'] == 1


def test_R3_unsafe_overlapping_paths_still_reject_globally(tmp_path):
    from mp4_analysis.thin.batch_cli import load_plan
    plan = tmp_path/'plan.json'
    write_json(plan, {'schema':1, 'output_root':'source/out', 'documents':[
        {'id':'one','source':'source','full_image':True}]})
    with pytest.raises(ValueError, match='disjoint'):
        load_plan(plan)


def test_R5_standalone_build_cannot_write_in_source(tmp_path):
    source = tmp_path/'source'
    source.mkdir()
    project = tmp_path/'project'
    batch = project/'batches'/'b_test'
    run = batch/'runs'/'r_test'
    run.mkdir(parents=True)
    write_json(run/'report.json', {})
    write_json(run/'crop_approval.json', {})
    write_json(batch/'manifest.json', {'source_directory': str(source)})
    write_json(project/'.screenshot_project.json', {'source_directory':str(source)})
    with pytest.raises(ValueError, match='disjoint'):
        build_document(run, source/'derived'/'v1', word=False, xlsx=False)
    assert not (source/'derived').exists()


def test_R5_standalone_build_protects_whole_ocr_project(tmp_path):
    run = tmp_path/'project'/'batches'/'b_test'/'runs'/'r_test'
    with pytest.raises(ValueError, match='disjoint'):
        build_document(run, tmp_path/'project'/'outside_batch'/'delivery', word=False, xlsx=False)
    assert not (tmp_path/'project').exists()


def test_R5_collection_delivery_sibling_remains_legal(tmp_path):
    doc = {'source_protection': {'source_directory':str(tmp_path/'source'),
                               'ocr_project_directory':str(tmp_path/'collection'/'d_one')}}
    target = tmp_path/'collection'/'deliveries'/'d_one'/'v1'
    assert require_delivery_target(doc, target) == target
    assert not target.exists()


def test_R6_single_prepare_is_in_elapsed(monkeypatch, tmp_path):
    from mp4_analysis.thin import pipeline
    clock = [10.0]
    monkeypatch.setattr(pipeline, 'time', SimpleNamespace(monotonic=lambda:clock[0]))
    def prepare(*args, **kwargs):
        clock[0] += 5.0
        return {'frames':[], 'approval_id':'test'}
    monkeypatch.setattr(pipeline, 'prepare_screenshots', prepare)
    result = pipeline.run(tmp_path/'source', tmp_path/'output', full_image=True, prepare_only=True)
    assert result['prepare_seconds'] == result['elapsed_seconds'] == 5.0
    assert result['inference_performed'] is False


def test_R6_collection_prepare_is_in_elapsed(monkeypatch, tmp_path):
    from mp4_analysis.thin import batch_cli
    clock = [10.0]
    monkeypatch.setattr(batch_cli, 'time', SimpleNamespace(monotonic=lambda:clock[0]))
    def prepare(plan, root):
        clock[0] += 7.0
        return {'approval_id':'same-identity', 'documents':[]}
    monkeypatch.setattr(batch_cli, '_prepare', prepare)
    plan = {'output_root':str(tmp_path/'collection'),'plan_file':str(tmp_path/'plan.json')}
    result = batch_cli.execute(plan)
    assert result['prepare_seconds'] == result['elapsed_seconds'] == 7.0
    assert result['approval_id'] == 'same-identity'
