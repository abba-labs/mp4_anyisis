"""Compare re-export tables to unchanged native XLSX, by exact native HTML identity.

This is exporter consistency, not source-video cell correctness. Never renumber
columns or alter text to obtain agreement; whitespace is the only normalization.
"""
import argparse
import json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from mp4_analysis.thin.video import file_hash


def snapshot(table):
    groups = {}
    for r, row in enumerate(table.rows, 1):
        for c, cell in enumerate(row.cells, 1):
            entry = groups.setdefault(cell._tc, {'row': r, 'column': c, 'last_row': r,
                'last_column': c, 'text': ''.join(cell.text.split())})
            entry.update(last_row=max(entry['last_row'], r), last_column=max(entry['last_column'], c))
    return list(groups.values())


def check(native, derived):
    records = []
    for cache_path in sorted(native.glob('*/adapter.json')):
        job = cache_path.parent
        candidate = derived/job.name/'document.docx'
        if not candidate.exists():
            records.append({'job': job.name, 'status': 'NOT_EXPORTED'})
            continue
        raw = json.loads(next(job.glob('*_res.json')).read_text(encoding='utf-8'))
        raw = raw.get('res', raw)
        blocks = [b for b in raw['parsing_res_list'] if b['block_label'] == 'table']
        htmls = {h.read_text(encoding='utf-8'): h for h in job.glob('*.html')}
        tables = Document(candidate).tables
        original = Document(next(job.glob('*.docx'))).tables
        if len(blocks) != len(tables):
            records.append({'job': job.name, 'status': 'TABLE_COUNT_MISMATCH'})
            continue
        for i, block in enumerate(blocks):
            html = htmls.get(block['block_content'])
            if html is None:
                records.append({'job': job.name, 'table': i, 'status': 'SOURCE_MAPPING_UNRESOLVED'})
                continue
            ws = load_workbook(html.with_suffix('.xlsx')).active
            merged = {m.start_cell.coordinate: m for m in ws.merged_cells.ranges}
            expected = []
            for row in ws:
                for cell in row:
                    if isinstance(cell, MergedCell): continue
                    span = merged.get(cell.coordinate)
                    text = '' if cell.value is None else str(cell.value)
                    expected.append({'row':cell.row, 'column':cell.column,
                        'last_row':span.max_row if span else cell.row,
                        'last_column':span.max_col if span else cell.column,
                        'text':''.join(text.split())})
            actual = snapshot(tables[i])
            old = snapshot(original[i]) if i < len(original) else None
            records.append({'job': job.name, 'table':i, 'source_bbox':block.get('block_bbox'),
                'native_html':html.name, 'native_html_sha256':file_hash(html),
                'status':'MATCH_NATIVE_XLSX' if actual == expected else 'MISMATCH',
                'old_matches_native_xlsx':old == expected, 'expected':expected, 'derived':actual})
    return records


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('native',type=Path)
    parser.add_argument('derived',type=Path)
    parser.add_argument('-o','--output',type=Path,required=True)
    args = parser.parse_args()
    records = check(args.native,args.derived)
    args.output.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
    from collections import Counter
    print(dict(Counter(r['status'] for r in records)))
